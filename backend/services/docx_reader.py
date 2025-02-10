"""
Module for reading and parsing DOCX files.
"""
from typing import Dict, Any, List, Optional, Union
from docx import Document
from werkzeug.datastructures import FileStorage
import logging
from pathlib import Path
import tempfile
import os

logger = logging.getLogger(__name__)

class DocxReader:
    def __init__(self):
        self.document = None
        self.temp_files: List[str] = []

    def __del__(self):
        """Cleanup temporary files on object destruction."""
        for temp_file in self.temp_files:
            try:
                if os.path.exists(temp_file):
                    os.unlink(temp_file)
            except Exception as e:
                logger.warning(f"Failed to delete temporary file {temp_file}: {e}")

    def _save_temp_file(self, file_storage: FileStorage) -> str:
        """Save FileStorage object to a temporary file and return its path."""
        temp_fd, temp_path = tempfile.mkstemp(suffix='.docx')
        try:
            os.close(temp_fd)
            file_storage.save(temp_path)
            self.temp_files.append(temp_path)
            return temp_path
        except Exception as e:
            if os.path.exists(temp_path):
                os.unlink(temp_path)
            raise ValueError(f"Failed to save temporary file: {str(e)}")

    def load_document(self, file_input: Union[str, FileStorage, Path]) -> None:
        """
        Load a DOCX document from various input types.
        
        Args:
            file_input: Can be a string path, FileStorage object, or Path object
        """
        try:
            if isinstance(file_input, FileStorage):
                file_path = self._save_temp_file(file_input)
            else:
                file_path = str(file_input)
                if not os.path.exists(file_path):
                    raise FileNotFoundError(f"File not found: {file_path}")

            self.document = Document(file_path)
            logger.info(f"Successfully loaded document from {file_path}")
        except Exception as e:
            logger.error(f"Error loading document: {str(e)}")
            raise

    def extract_text_by_paragraphs(self) -> List[str]:
        """Extract text from paragraphs in the document."""
        if not self.document:
            raise ValueError("No document loaded")
        
        return [paragraph.text for paragraph in self.document.paragraphs if paragraph.text.strip()]

    def extract_tables(self) -> List[List[List[str]]]:
        """Extract tables from the document."""
        if not self.document:
            raise ValueError("No document loaded")
        
        tables = []
        for table in self.document.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                if any(row_data):  # Only include rows that have some content
                    table_data.append(row_data)
            if table_data:  # Only include tables that have some content
                tables.append(table_data)
        return tables

    def extract_structured_content(self) -> Dict[str, Any]:
        """
        Extract content from the document in a structured format.
        Returns a dictionary with paragraphs and tables.
        """
        if not self.document:
            raise ValueError("No document loaded")

        return {
            "paragraphs": self.extract_text_by_paragraphs(),
            "tables": self.extract_tables()
        }

    def find_field_content(self, field_marker: str) -> Optional[str]:
        """
        Find content associated with a specific field marker in the document.
        
        Args:
            field_marker: The text marker that indicates the field (e.g., "[Field 1]")
            
        Returns:
            The content following the field marker, or None if not found
        """
        if not self.document:
            raise ValueError("No document loaded")

        # Search in paragraphs
        for i, para in enumerate(self.document.paragraphs):
            if field_marker in para.text:
                # If there's a next paragraph, return its content
                if i + 1 < len(self.document.paragraphs):
                    return self.document.paragraphs[i + 1].text.strip()
                break

        # Search in tables
        for table in self.document.tables:
            for i, row in enumerate(table.rows):
                for j, cell in enumerate(row.cells):
                    if field_marker in cell.text:
                        # Try to get content from next cell or row
                        if j + 1 < len(row.cells):
                            return row.cells[j + 1].text.strip()
                        elif i + 1 < len(table.rows):
                            return table.rows[i + 1].cells[j].text.strip()

        return None
