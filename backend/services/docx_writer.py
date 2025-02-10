"""
Module for writing formatted data to DOCX files.
"""
from typing import Dict, Any
from docx import Document
from docx.shared import Inches
import logging

logger = logging.getLogger(__name__)

class DocxWriter:
    def __init__(self, template_path: str):
        self.template_path = template_path
        self.document = Document(template_path)

    def write_field(self, field_name: str, content: str) -> None:
        """Write content to a specific field in the template."""
        try:
            # Find and replace the field placeholder in the document
            for paragraph in self.document.paragraphs:
                if field_name in paragraph.text:
                    paragraph.text = paragraph.text.replace(field_name, content)
                    
            # Also check tables
            for table in self.document.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if field_name in cell.text:
                            cell.text = cell.text.replace(field_name, content)
        except Exception as e:
            logger.error(f"Error writing field {field_name}: {str(e)}")
            raise

    def save(self, output_path: str) -> None:
        """Save the document to the specified path."""
        try:
            self.document.save(output_path)
            logger.info(f"Document saved successfully to {output_path}")
        except Exception as e:
            logger.error(f"Error saving document: {str(e)}")
            raise
